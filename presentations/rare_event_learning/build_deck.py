"""Build the rare-event teaching deck. Run with python-pptx, matplotlib and Pillow.
All numeric examples are analytical illustrations, not measured RESuM results.
"""
from pathlib import Path
import io
import textwrap
import numpy as np
import matplotlib
matplotlib.use('Agg')
import matplotlib.pyplot as plt
from matplotlib.backends.backend_pdf import PdfPages
from matplotlib.patches import Rectangle
from matplotlib import font_manager
from PIL import Image, ImageFont
from pptx import Presentation
from pptx.util import Inches, Pt
from pptx.dml.color import RGBColor
from pptx.enum.shapes import MSO_SHAPE

ROOT = Path(__file__).resolve().parent
OUT = ROOT / 'rare_event_learning'
W,H = 13.333333,7.5
NAVY='#17324D'; TEAL='#008B8B'; ORANGE='#D46C37'; GRAY='#576675'; LIGHT='#EDF3F6'; BG='#FAFCFD'
plt.rcParams.update({'font.family':'DejaVu Sans','font.size':13,'axes.spines.top':False,'axes.spines.right':False,'axes.labelcolor':NAVY,'text.color':NAVY})
prs=Presentation(); prs.slide_width=Inches(W); prs.slide_height=Inches(H)
pdf=PdfPages(str(OUT)+'.pdf'); previews=[]
FONT=font_manager.findfont('DejaVu Sans'); BOLD=font_manager.findfont(font_manager.FontProperties(family='DejaVu Sans',weight='bold'))

def rgb(c): return RGBColor.from_string(c.lstrip('#'))

def wrap(s,w,size,bold=False):
    font=ImageFont.truetype(BOLD if bold else FONT,round(size*4))
    limit=w*72*4
    out=[]
    for line in s.split('\n'):
        current=''
        for word in line.split():
            candidate=(current+' '+word).strip()
            if font.getlength(candidate)>limit and current: out.append(current); current=word
            else: current=candidate
        out.append(current)
    return '\n'.join(out)

class Slide:
    def __init__(self,title,kicker='RARE EVENTS • OPTICAL MAP',notes=''):
        self.s=prs.slides.add_slide(prs.slide_layouts[6]); self.s.background.fill.solid(); self.s.background.fill.fore_color.rgb=rgb(BG)
        self.fig=plt.figure(figsize=(W,H)); self.fig.patch.set_facecolor(BG)
        self.ax=self.fig.add_axes([0,0,1,1]); self.ax.set(xlim=(0,W),ylim=(H,0)); self.ax.axis('off')
        self.text(.55,.28,12,.25,kicker,10,TEAL,True)
        self.text(.55,.78,12.25,1.0,title,29,NAVY,True)
        self.text(.55,7.13,11.5,.2,'Illustrative examples • No experimental model ranking',9,GRAY)
        self.text(12.3,7.1,.5,.25,str(len(prs.slides)),11,GRAY)
        self.s.notes_slide.notes_text_frame.text=notes or 'Numerical examples are analytical illustrations; they are not outcomes from the optical-map experiments.'
    def text(self,x,y,w,h,s,size=21,color=NAVY,bold=False):
        while True:
            lines=wrap(s,w,size,bold)
            if len(lines.split('\n'))*size*1.23/72 <= h or size<=11: break
            size-=1
        if len(lines.split('\n'))*size*1.23/72>h+.02: raise ValueError(('text overflow',s))
        self.ax.text(x,y,lines,fontsize=size,color=color,weight='bold' if bold else 'normal',va='top',linespacing=1.23)
        tb=self.s.shapes.add_textbox(Inches(x), Inches(y), Inches(w), Inches(h)); tf=tb.text_frame
        tf.clear(); tf.word_wrap=False; tf.margin_left=tf.margin_right=tf.margin_top=tf.margin_bottom=0
        for i,line in enumerate(lines.split('\n')):
            p=tf.paragraphs[0] if i==0 else tf.add_paragraph(); p.text=line;p.font.name='DejaVu Sans';p.font.size=Pt(size);p.font.bold=bold;p.font.color.rgb=rgb(color);p.space_after=Pt(0);p.line_spacing=1.23
    def box(self,x,y,w,h,color=LIGHT):
        self.ax.add_patch(Rectangle((x,y),w,h,color=color,zorder=0))
        sh=self.s.shapes.add_shape(MSO_SHAPE.RECTANGLE, Inches(x),Inches(y),Inches(w),Inches(h));sh.fill.solid();sh.fill.fore_color.rgb=rgb(color);sh.line.fill.background()
    def card(self,x,y,w,h,heading,body,color=TEAL):
        self.box(x,y,w,h);self.text(x+.22,y+.2,w-.44,.62,heading,23,color,True);self.text(x+.22,y+1,w-.44,h-1.15,body,21)
    def takeaway(self,s):
        self.box(.55,6.28,12.23,.58,NAVY);self.text(.75,6.4,11.83,.36,s,17,'#FFFFFF',True)
    def chart(self,fig,x,y,w,h):
        b=io.BytesIO();fig.savefig(b,format='png',dpi=170,bbox_inches='tight',facecolor=BG);b.seek(0)
        self.ax.imshow(Image.open(b),extent=(x,x+w,y+h,y),aspect='auto',zorder=2)
        b.seek(0);self.s.shapes.add_picture(b,Inches(x),Inches(y),width=Inches(w),height=Inches(h));plt.close(fig)
    def save(self):
        pdf.savefig(self.fig)
        dest=ROOT/f'preview_{len(prs.slides):02}.png';self.fig.savefig(dest,dpi=85);previews.append(dest);plt.close(self.fig)

s=Slide('Why rare detections are hard to learn')
s.text(.65,1.8,7.2,1.3,'Only 5 detections in 10,000 events',34,NAVY,True)
s.text(.65,3.25,6.9,1.6,'At a 0.05% detection rate, a network can look excellent while missing every positive event.',25)
s.card(8.2,1.9,4.3,3.8,'Three different problems','Misleading accuracy\n\nSparse evidence\n\nShifted probabilities')
s.takeaway('The goal: learn rare-event structure without distorting the population rate.');s.save()

s=Slide('99.95% accuracy — and zero detections found')
f,ax=plt.subplots(figsize=(7,3.8));idx=np.arange(10000);ax.scatter(idx%200,idx//200,s=3,c='#BFCBD3',linewidths=0);pos=np.array([703,2808,4900,7255,9737]);ax.scatter(pos%200,pos//200,s=45,c=ORANGE,edgecolor='white');ax.set_title('10,000 events: 9,995 negatives + 5 positives');ax.axis('off');s.chart(f,.6,1.75,7.6,3.7)
s.card(8.5,1.9,4.1,3.8,'Predict “negative” for all','Accuracy: 99.95%\n\nPositive recall: 0%\n\nEvery detection is missed.',ORANGE)
s.takeaway('Accuracy counts successes mostly on the abundant negative class.');s.save()

s=Slide('A negative decision is not a zero probability',notes='For threshold 0.5, any q<0.5 produces a negative decision. Constant q=prevalence minimizes unweighted population BCE among constant predictors, but need not be optimal among feature-dependent predictors. Exact q=0 has infinite BCE if any positives occur.')
s.card(.65,1.85,5.85,3.9,'Classification decision','“Does this event produce a hit?”\n\nIf the threshold is 50%, a score of 0.05% is classified as negative.')
s.card(6.8,1.85,5.85,3.9,'Probability estimate','“How likely is a hit?”\n\nA score of 0.05% for all events predicts 5 hits per 10,000 on average.')
s.takeaway('A model can have zero recall at threshold 0.5 and still estimate the mean rate.');s.save()

s=Slide('BCE does not reward predicting exactly zero',notes='Analytical constant-predictor loss L(q)=-pi log(q)-(1-pi)log(1-q), pi=0.0005. Its unique minimum is q=pi. BCE gradient with respect to logit is q-y. At q=0 positive-event loss diverges. Numerically stable BCE with logits avoids explicit log(0); it does not rebalance classes.')
f,ax=plt.subplots(figsize=(6.8,3.9));q=np.logspace(-7,-1,600);pi=.0005;loss=-pi*np.log(q)-(1-pi)*np.log1p(-q);ax.semilogx(q*100,loss,color=TEAL,lw=2.7);ax.axvline(.05,c=ORANGE,ls='--');ax.set(xlabel='Constant predicted probability (%)',ylabel='Average binary cross-entropy',ylim=(0,.014));ax.annotate('Best constant prediction\nq = 0.05%',xy=(.05,-pi*np.log(pi)-(1-pi)*np.log1p(-pi)),xytext=(.00002,.011),arrowprops={'arrowstyle':'->','color':NAVY});s.chart(f,.6,1.7,7,4.25)
s.text(8,1.95,4.65,3.9,'A positive event has loss −log(q).\n\nAs q approaches zero, that penalty grows without bound.\n\nThe best constant score is the true positive fraction.',23)
s.takeaway('Near-zero scores can be sensible; failing to learn useful variation is the concern.');s.save()

s=Slide('Many small batches contain no positive example',notes='Independent natural draws with prevalence pi=0.0005: P(no positives)=(1-pi)^B. Correlated/grouped sampling can change these numbers. Repeated epochs revisit the same dataset; they do not supply new independent positives.')
f,ax=plt.subplots(figsize=(6.5,3.7));bs=np.array([128,512,2048]);v=100*(1-.0005)**bs;ax.bar([0,1,2],v,color=[ORANGE,ORANGE,TEAL],width=.55);ax.set(xticks=[0,1,2],xticklabels=bs,xlabel='Events per batch',ylabel='Batches with no positives (%)',ylim=(0,110));[ax.text(i,a+2,f'{a:.1f}%',ha='center',fontsize=17) for i,a in enumerate(v)];s.chart(f,.65,1.85,7,3.95)
s.text(8,1.9,4.6,3.8,'At 0.05% positives:\n\nA batch of 512 has only 0.256 positives on average.\n\nMost such updates contain only negative examples.',23)
s.takeaway('Larger or stratified batches improve exposure; they do not create new evidence.');s.save()

s=Slide('Early training can mostly learn “make scores smaller”',notes='For unweighted mean BCE with sigmoid logits z, dL/dz=q-y per example before averaging. Illustration assumes all q=0.5 and examines a shared output bias: 9995*0.5=4997.5, 5*(-0.5)=-2.5. Parameter gradients also depend on features/Jacobians. At q=pi=0.0005 the shared-bias contributions balance at +4.9975 and -4.9975; class counts alone do not imply permanent gradient domination.')
s.card(.65,1.85,5.85,3.75,'9,995 negative events','Start with scores near 50%.\n\nTogether, negatives strongly push the shared output bias downward.',ORANGE)
s.card(6.8,1.85,5.85,3.75,'5 positive events','Each positive pushes upward.\n\nThere are very few examples from which to learn a generalizable positive pattern.')
s.text(.85,5.75,11.8,.4,'As scores fall, easy-negative gradients shrink. This is not an unavoidable zero-loss trap.',18)
s.takeaway('Learning the overall rarity can be much easier than learning who is more likely.');s.save()

s=Slide('Oversampling makes positives visible — not more diverse',notes='Random positive oversampling repeats existing observations. It can improve optimizer exposure to minority patterns but does not increase the independent positive sample size. Risks described are conditional, not claims that oversampling always fails.')
s.card(.65,1.9,3.85,3.85,'Original data','5 distinct positives\n\nA few examples of what a detection looks like.')
s.card(4.75,1.9,3.85,3.85,'Oversampled batches','The same 5 positives appear repeatedly.\n\nMore training exposure.')
s.card(8.85,1.9,3.85,3.85,'Possible failure','The model memorizes those examples.\n\nUnseen positives remain difficult.',ORANGE)
s.takeaway('More repetitions help optimization; more independent positives help generalization.');s.save()

s=Slide('Changing the class balance changes what scores mean',notes='Toy population with no informative features: natural prevalence pi=0.0005. BCE-trained constant predictors on sampled prevalences r=0.05 or 0.25 minimize loss at r, not pi. In a feature-dependent model pure class-dependent resampling shifts posterior odds. This does not imply every real trained model has a uniform 100x or 500x error.')
s.text(.8,1.85,11.8,.65,'Toy example: the features contain no information about the label.',23)
for x,title,num,body,c in [( .7,'Natural data','0.05%','Best constant BCE score',TEAL),(4.9,'5% positive batches','5%','100× the population rate',ORANGE),(9.1,'25% positive batches','25%','500× the population rate',ORANGE)]:
    s.box(x,2.7,3.55,2.85);s.text(x+.18,2.9,3.19,.7,title,22,c,True);s.text(x+.18,3.75,3.19,.9,num,40,c,True);s.text(x+.18,4.9,3.19,.45,body,16)
s.takeaway('Better positive recall can come with an upward bias in predicted detection rates.');s.save()

s=Slide('Reweighting also changes the training objective',notes='Weighted BCE L=-w*y*log(q)-(1-y)*log(1-q), with constant positive weight w and negative weight 1. At true conditional probability p, optimum q=w*p/(1-p+w*p). At p=0.0005, w=100 gives 4.764%; w=1999 gives 50%. Correct inverse-sampling weights are different: they restore the natural expected loss. Large weights can increase stochastic gradient variance when positives are rare.')
s.card(.65,1.9,5.85,3.9,'What it does','A positive mistake receives a larger penalty.\n\nThis can encourage the network to learn rare positive patterns.')
s.card(6.8,1.9,5.85,3.9,'Why it can disappoint','A few positives can drive noisy updates or overfitting.\n\nThe score is no longer automatically a natural detection probability.',ORANGE)
s.takeaway('Weighting changes importance; it cannot add missing information.');s.save()

s=Slide('Focal loss can focus on ambiguity as well as signal',notes='Reference: Lin et al., Focal Loss for Dense Object Detection, ICCV 2017, https://openaccess.thecvf.com/content_iccv_2017/html/Lin_Focal_Loss_for_ICCV_2017_paper.html . Focal loss downweights well-classified examples using (1-p_t)^gamma. Hard examples may include informative positives, ambiguous events or label errors. In inherently stochastic detection the irreducible component cannot be learned away. Risks are task-specific inferences, not experimental results from that paper.')
s.card(.65,1.9,5.85,3.8,'Useful when…','Many easy negatives consume training attention.\n\nReducing their influence makes harder examples matter more.')
s.card(6.8,1.9,5.85,3.8,'Can fail when…','“Hard” means noisy or intrinsically unpredictable.\n\nThe model spends capacity on examples with little learnable signal.',ORANGE)
s.text(.85,5.85,11.6,.3,'Source: Lin et al. (2017). Optical-map implications are hypotheses to test.',13,GRAY)
s.takeaway('A different loss cannot turn random outcomes into predictable events.');s.save()

s=Slide('Mixup adds interpolation — not independent detections',notes='Reference: Zhang et al., mixup: Beyond Empirical Risk Minimization, ICLR 2018, https://arxiv.org/abs/1710.09412 . x_tilde=lambda*x_plus+(1-lambda)*x_minus and y_tilde=lambda. If every pair is positive-negative and E[lambda]=0.5, mean synthetic label is 0.5. Natural independent pairs and label-independent lambda preserve prevalence in expectation. The derivation is our illustration; claims about unphysical interpolation are conditional on features/task.')
s.text(.8,1.8,11.8,.7,'Positive (label 1) + negative (label 0) → synthetic event (label λ)',24,TEAL,True)
s.card(.65,2.8,5.85,2.9,'Why it can help','Smoother predictions between examples; less memorization of individual training events.')
s.card(6.8,2.8,5.85,2.9,'Why it can fail','Interpolated features may not represent real events. Forced positive–negative pairing changes the label distribution.',ORANGE)
s.text(.85,5.8,11.6,.35,'If every pair is +/− and mean λ = 0.5, the synthetic mean label is 50%.',18)
s.takeaway('Natural pairing preserves the mean in expectation, but rarely includes positives.');s.save()

s=Slide('Some events may be impossible to classify reliably',notes='Conceptual distinction between Bayes uncertainty and insufficient data/model capacity. If relevant randomness is unobserved, P(Y=1|X) can remain small for all X. Balanced accuracy or rare-event recall may improve with an altered decision threshold; perfect classification is not implied by learning calibrated probabilities. This is a possible explanation, not a diagnosis of the optical dataset.')
s.card(.65,1.9,5.85,3.85,'Learnable structure','Position or event features change the chance of detection.\n\nThe network can learn which events are more likely.')
s.card(6.8,1.9,5.85,3.85,'Unobserved randomness','Similar recorded inputs can still produce different outcomes.\n\nNo loss or sampler can reveal information absent from the inputs.',ORANGE)
s.takeaway('Accurate probabilities can be achievable even when individual hit prediction is weak.');s.save()

s=Slide('For the optical map, protect the population mean',notes='Recommendations, not completed experiments. Evaluate on untouched natural prevalence, with spatially grouped train/validation/test splits. A separate real-data BCE plus weighted synthetic auxiliary loss anchors training but does not guarantee calibration. Mean bias is predicted minus observed fraction, preferably accompanied by MAE/trend and uncertainty. Average precision baseline equals prevalence for random ranking in the population limit.')
s.card(.65,1.9,3.85,3.95,'Training','Improve exposure to positives.\n\nRetain a loss on naturally sampled real events.\n\nTune augmentation strength.')
s.card(4.75,1.9,3.85,3.95,'Validation','Use the natural class balance.\n\nHold out whole voxels.\n\nCompare several seeds.')
s.card(8.85,1.9,3.85,3.95,'Measure both','Event ranking:\nprecision–recall / AP\n\nPopulation map:\nmean bias, MAE and spatial trend')
s.takeaway('Success means useful ranking AND trustworthy rates; accuracy alone proves neither.');s.save()

s=Slide('What would most directly improve the evidence?',notes='Analytical homogeneous-rate illustration pi=0.0005. Expected positives N*pi; P(zero)=(1-pi)^N. Values are not claims about every optical voxel, whose rates vary. Increasing primaries improves count precision; adding spatial locations improves coverage. The allocation tradeoff is an experimental design question.')
s.text(.8,1.85,11.7,.6,'At a true detection probability of 0.05%:',25)
for x,n in [(.7,1500),(4.9,5000),(9.1,20000)]:
    s.box(x,2.65,3.55,2.9);s.text(x+.2,2.9,3.15,.7,f'{n:,} primaries',23,TEAL,True);s.text(x+.2,3.8,3.15,1.5,f'{n*.0005:g} expected hits\n\n' + (f'{100*(1-.0005)**n:.1f}%' if 100*(1-.0005)**n >= .01 else '<0.01%') + ' chance of zero hits',21)
s.text(.85,5.75,11.8,.4,'More primaries reveal rare detections; more voxel locations reveal spatial structure.',19)
s.takeaway('Sampling tricks redistribute the evidence. Simulation supplies new evidence.');s.save()

s=Slide('Appendix: why weights and sampling bias probabilities',kicker='TECHNICAL APPENDIX',notes='Analytical derivations: weighted BCE optimum found by setting derivative to zero. Prior-shift correction follows Bayes odds and requires class-only sampling with unchanged class-conditional features, a calibrated sampled posterior q and known natural prevalence pi. It is not generally valid for arbitrary mixup or within-voxel selection changes. Related reading: Menon et al., Long-tail learning via logit adjustment, ICLR 2021, https://research.google/pubs/long-tail-learning-via-logit-adjustment/ .')
s.text(.75,1.75,12,1.0,'Positive weight w changes the BCE optimum:\nq* = w p / (1 − p + w p)',25,TEAL,True)
s.text(.75,3.05,12,1,'At p = 0.05%: w = 100 gives q* ≈ 4.76%; w = 1,999 gives q* = 50%.',22)
s.text(.75,4.1,12,1,'Under pure class-dependent resampling:\nlogit(p) = logit(q) + logit(π) − logit(r)',24,TEAL,True)
s.text(.75,5.25,12,.8,'π: natural positive fraction; r: sampled fraction; q: calibrated sampled score.\nRequires unchanged within-class distributions; not a general mixup correction.',19)
s.takeaway('Correct sampling weights can restore the original objective; arbitrary class weights do not.');s.save()

s=Slide('References and scope',kicker='SOURCES & SPEAKER NOTES')
s.text(.75,1.75,11.9,1.05,'Lin et al. (2017) — Focal Loss for Dense Object Detection\nopenaccess.thecvf.com/content_iccv_2017/html/Lin_Focal_Loss_for_ICCV_2017_paper.html',18)
s.text(.75,3.0,11.9,.9,'Zhang et al. (2018) — mixup: Beyond Empirical Risk Minimization\nhttps://arxiv.org/abs/1710.09412',18)
s.text(.75,4.1,11.9,1.05,'Menon et al. (2021) — Long-tail learning via logit adjustment\nhttps://research.google/pubs/long-tail-learning-via-logit-adjustment/',18)
s.text(.75,5.35,11.9,.65,'All numbers and diagrams here are illustrative calculations, not new training results.\nFailure modes describe possibilities, not inevitable outcomes of these methods.',18,GRAY)
s.takeaway('Detailed assumptions and derivations are included in the PowerPoint speaker notes.');s.save()

pdf.close();prs.save(str(OUT)+'.pptx')
# A compact contact sheet makes visual review straightforward.
thumbs=[]
for p in previews:
    im=Image.open(p).convert('RGB');im.thumbnail((480,270));thumbs.append(im.copy())
sheet=Image.new('RGB',(480*4,270*((len(thumbs)+3)//4)), 'white')
for i,im in enumerate(thumbs):sheet.paste(im,((i%4)*480,(i//4)*270))
sheet.save(ROOT/'contact_sheet.png')
print(f'Built {len(prs.slides)} slides: {OUT}.pptx / .pdf')
